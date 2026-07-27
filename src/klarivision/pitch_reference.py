"""Turkish music pitch reference loaded from the musician-maintained workbook."""

from __future__ import annotations

from dataclasses import dataclass, replace
from pathlib import Path
import re

from openpyxl import load_workbook

from .runtime_paths import resource_root


PROJECT_ROOT = resource_root()
DEFAULT_REFERENCE_PATH = (
    PROJECT_ROOT / "data" / "reference" / "Turk_Muzigi_Perdeleri_ve_Mikrotonal_Notasyon.xlsx"
)


@dataclass(frozen=True)
class TurkishPitchReference:
    """One AEU Turkish-music pitch line for Sol clarinet notation."""

    frequency_hz: float
    turkish_name: str
    sol_clarinet_note: str
    piano_note: str
    koma_description: str
    octave_offset: int = 0

    @property
    def accidental_notation(self) -> str:
        """The source notation before its explanatory parenthesis."""
        return self.koma_description.split(" (", maxsplit=1)[0].strip()

    @property
    def display_notation(self) -> str:
        """Portable Sol-clarinet label such as ``Re ♭5`` or ``Fa ♯1``.

        The workbook remains the source of truth.  This display-only form
        avoids relying on rarely installed AEU music-symbol fonts while
        retaining the direction and exact comma amount.
        """
        note_match = re.match(r"^(Do|Re|Mi|Fa|Sol|La|Si)\b", self.accidental_notation)
        koma_match = re.search(r"(\d+)\s*Koma", self.koma_description, re.IGNORECASE)
        if not note_match or not koma_match:
            return self.accidental_notation

        source_symbol = self.accidental_notation[len(note_match.group(0)) :]
        if any(symbol in source_symbol for symbol in ("♯", "𝄰", "𝄱", "𝄵")):
            accidental = "♯"
        elif any(symbol in source_symbol for symbol in ("♭", "𝄳", "𝄴")):
            accidental = "♭"
        elif "Diyez" in self.koma_description:
            accidental = "♯"
        elif "Bemol" in self.koma_description:
            accidental = "♭"
        else:
            return self.accidental_notation
        return f"{note_match.group(0)} {accidental}{koma_match.group(1)}"

    @property
    def octave_label(self) -> str:
        if self.octave_offset < 0:
            return "alt oktav"
        if self.octave_offset > 0:
            return "üst oktav"
        return "referans oktav"


def load_turkish_pitch_reference(
    path: Path = DEFAULT_REFERENCE_PATH,
) -> tuple[TurkishPitchReference, ...]:
    """Read the user-maintained Turkish microtonal notation workbook."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook["Türk Müziği Perdeleri"]
        header_row = next(
            (
                index
                for index, row in enumerate(sheet.iter_rows(values_only=True), start=1)
                if row and row[0] == "Frekans (Hz)"
            ),
            None,
        )
        if header_row is None:
            raise ValueError("Referans dosyasında 'Frekans (Hz)' başlık satırı bulunamadı.")
        headers = [
            str(value).strip() if value is not None else ""
            for value in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
        ]
        positions = {name: index for index, name in enumerate(headers)}
        required = {
            "Frekans (Hz)",
            "Türk Nota İsmi",
            "TM Notası (Sol Klarnet)",
            "Batı Notası(Piyano)",
            "Koma Açıklaması",
        }
        missing = required - positions.keys()
        if missing:
            raise ValueError(f"Referans dosyasında eksik sütunlar var: {', '.join(sorted(missing))}")

        records: list[TurkishPitchReference] = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            frequency = row[positions["Frekans (Hz)"]]
            if frequency is None:
                continue
            records.append(
                TurkishPitchReference(
                    frequency_hz=float(frequency),
                    turkish_name=str(row[positions["Türk Nota İsmi"]] or "").strip(),
                    sol_clarinet_note=str(row[positions["TM Notası (Sol Klarnet)"]] or "").strip(),
                    piano_note=str(row[positions["Batı Notası(Piyano)"]] or "").strip(),
                    koma_description=str(row[positions["Koma Açıklaması"]] or "").strip(),
                )
            )
    finally:
        workbook.close()

    if len(records) < 2:
        raise ValueError("Referans dosyasında yeterli perde satırı bulunamadı.")
    return tuple(records)


def extend_reference_octaves(
    records: tuple[TurkishPitchReference, ...],
    *,
    lower_octaves: int = 1,
    upper_octaves: int = 1,
) -> tuple[TurkishPitchReference, ...]:
    """Extend the source octave by exact ×½/×2 frequency relationships.

    Generated rows preserve the source's Sol-clarinet notation and AEU comma
    symbols.  Their Turkish name intentionally retains the source name plus an
    octave label rather than inventing a musicological name.
    """
    ordered = tuple(sorted(records, key=lambda record: record.frequency_hz))
    body, endpoint = ordered[:-1], ordered[-1]
    extended: list[TurkishPitchReference] = []
    for offset in range(-lower_octaves, upper_octaves + 1):
        extended.extend(
            replace(record, frequency_hz=record.frequency_hz * 2**offset, octave_offset=offset)
            for record in body
        )
    extended.append(
        replace(endpoint, frequency_hz=endpoint.frequency_hz * 2**upper_octaves, octave_offset=upper_octaves)
    )
    return tuple(extended)

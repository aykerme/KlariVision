"""Human-readable quality-control report for the editable pitch workbook."""

from __future__ import annotations

import html
import math
from pathlib import Path

from openpyxl import load_workbook


def _nearest_equal_temperament_note(frequency_hz: float) -> tuple[str, float]:
    """Return an approximate piano name only as a numeric cross-check."""
    midi = round(69 + 12 * math.log2(frequency_hz / 440))
    names = ("Do", "Do♯", "Re", "Re♯", "Mi", "Fa", "Fa♯", "Sol", "Sol♯", "La", "La♯", "Si")
    note = f"{names[midi % 12]}{midi // 12 - 1}"
    reference_hz = 440 * 2 ** ((midi - 69) / 12)
    cents = 1200 * math.log2(frequency_hz / reference_hz)
    return note, cents


def build_pitch_reference_report(workbook_path: Path, output_path: Path) -> None:
    """Render the workbook data as a read-only HTML review page.

    The report makes no claim that nearest equal-tempered notes are the final
    clarinet notation.  It merely exposes the current numbers in the editable
    workbook so that musical validation happens before UI integration.
    """
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        settings_sheet = workbook["Ayarlar"]
        mapping_sheet = workbook["Perde Eşleme"]
        settings = list(settings_sheet.iter_rows(min_row=3, max_row=7, values_only=True))
        headers = [str(value).strip() if value is not None else "" for value in next(mapping_sheet.iter_rows(min_row=2, max_row=2, values_only=True))]
        rows = [dict(zip(headers, values, strict=True)) for values in mapping_sheet.iter_rows(min_row=3, values_only=True) if values[0] is not None]
    finally:
        workbook.close()

    anchors = [row for row in rows if row["Türk müziği perdesi"] in {"Yegâh", "Dügâh", "Neva"}]
    status_counts: dict[str, int] = {}
    for row in rows:
        status = str(row.get("Onay") or "Belirtilmemiş")
        status_counts[status] = status_counts.get(status, 0) + 1

    def esc(value: object) -> str:
        return html.escape("" if value is None else str(value))

    anchor_cards = "".join(
        f"""<article class=anchor><h3>{esc(row['Türk müziği perdesi'])} ({esc(row['Parantez içi nota'])})</h3>
        <p><b>Duyulan:</b> {float(row['Duyulan Hz']):.2f} Hz</p>
        <p><b>Sol klarnet yazılı:</b> {float(row['Sol klarnet yazılı Hz']):.2f} Hz</p></article>"""
        for row in anchors
    )
    table_rows = "".join(
        f"""<tr><td>{int(row['Koma'])}</td><td>{esc(row['Türk müziği perdesi'])}</td>
        <td>{esc(row['Parantez içi nota'])}</td><td>{float(row['Duyulan Hz']):.2f}</td>
        <td>{float(row['Sol klarnet yazılı Hz']):.2f}</td>
        <td>{_nearest_equal_temperament_note(float(row['Sol klarnet yazılı Hz']))[0]}</td>
        <td>{_nearest_equal_temperament_note(float(row['Sol klarnet yazılı Hz']))[1]:+.1f} sent</td>
        <td>{esc(row.get('Onay'))}</td></tr>"""
        for row in rows
    )
    settings_rows = "".join(f"<tr><th>{esc(name)}</th><td>{esc(value)}</td></tr>" for name, value in settings)
    status_text = ", ".join(f"{esc(status)}: {count}" for status, count in sorted(status_counts.items()))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        f"""<!doctype html><html lang=tr><meta charset=utf-8><title>KlariVision · Perde eşleme kontrolü</title>
        <style>body{{font:15px system-ui,-apple-system,sans-serif;margin:0;background:#f5f7fa;color:#1e2936}}main{{max-width:1200px;margin:auto;padding:32px}}h1{{margin:0 0 8px}}p{{line-height:1.55}}.notice{{background:#fff3cd;border:1px solid #ebcf76;padding:14px 16px;border-radius:10px}}.cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(210px,1fr));gap:12px;margin:18px 0}}.anchor,.panel{{background:#fff;border:1px solid #dce3eb;border-radius:10px;padding:14px}}.anchor h3{{margin:0 0 10px}}.anchor p{{margin:4px 0}}table{{width:100%;border-collapse:collapse;background:#fff;font-size:13px}}th{{text-align:left;background:#eaf1f8}}th,td{{padding:8px;border-bottom:1px solid #e5e9ee;white-space:nowrap}}.scroll{{overflow:auto;border:1px solid #dce3eb;border-radius:10px}}.muted{{color:#596775;font-size:13px}}</style>
        <main><h1>KlariVision · Perde eşleme kontrol raporu</h1>
        <p class=muted>Kaynak: <code>{esc(workbook_path)}</code> · {len(rows)} perde satırı · Durumlar: {status_text}</p>
        <p class=notice>Bu sayfa yalnızca kaynak tablodaki sayıları gösterir; hiçbir eşlemeyi değiştirmez. “12 eşit aralıklı yakın nota” sütunu sadece frekans kontrolüdür, nihai Sol klarnet nota adı değildir.</p>
        <h2>Kontrol ankrajları</h2><section class=cards>{anchor_cards}</section>
        <section class=panel><h2>Ayarlar</h2><table>{settings_rows}</table></section>
        <h2>Tüm perde satırları</h2><div class=scroll><table><thead><tr><th>Koma</th><th>Türk müziği perdesi</th><th>Do–Re–Mi</th><th>Duyulan Hz</th><th>Sol klarnet yazılı Hz</th><th>12TET yakın nota</th><th>Fark</th><th>Onay</th></tr></thead><tbody>{table_rows}</tbody></table></div></main></html>""",
        encoding="utf-8",
    )
